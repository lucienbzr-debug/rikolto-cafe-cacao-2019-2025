# -*- coding: utf-8 -*-
"""Met en forme l'export Word du rapport (charte Rikolto : Arial, vert, tableaux à en-tête coloré)."""
import os, sys
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "rapport_export_source.docx")
OUT = os.path.join(HERE, "..", "..", "livrables", "Rapport_direction_Cafe_Cacao_2019-2025.docx")
GREEN, INK, MUTED, LINE = RGBColor(0x1F, 0x5C, 0x4A), RGBColor(0x1B, 0x26, 0x22), RGBColor(0x56, 0x63, 0x5C), "DCE1DA"

d = docx.Document(SRC)
def set_font(style, size=None, bold=None, color=None, name="Arial"):
    f = style.font; f.name = name
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"): rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if color is not None: f.color.rgb = color

st = d.styles
set_font(st["Normal"], 10, color=INK)
st["Normal"].paragraph_format.space_after = Pt(5); st["Normal"].paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 11)):
    if name in [s.name for s in st]:
        set_font(st[name], size, True, GREEN)
        pf = st[name].paragraph_format; pf.keep_with_next = True; pf.space_before = Pt(16 if name != "Heading 1" else 0); pf.space_after = Pt(6)

# texte courant : forcer Arial sur les runs qui portent une police explicite
for p in d.paragraphs:
    for r in p.runs:
        if r.font.name and r.font.name not in ("Arial",) and "Mono" not in (r.font.name or "") and "Courier" not in (r.font.name or ""):
            r.font.name = "Arial"

# ligne d'auteur en gris
if len(d.paragraphs) > 1:
    for r in d.paragraphs[1].runs: r.font.color.rgb = MUTED; r.font.size = Pt(9.5)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")): tcPr.remove(old)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor); tcPr.append(shd)

def borders(tbl, color=LINE):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")): tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color); b.append(e)
    tblPr.append(b)

for t in d.tables:
    borders(t)
    for i, row in enumerate(t.rows):
        if i == 0:
            trPr = row._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
        for cell in row.cells:
            if i == 0: shade(cell, "1F5C4A")
            elif i % 2 == 0: shade(cell, "F3F5F2")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.name = "Arial"; r.font.size = Pt(8.5)
                    if i == 0: r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# légendes des graphiques (paragraphe suivant une image) en gris italique
body = d.paragraphs
for i, p in enumerate(body[:-1]):
    if p._p.findall(".//" + qn("w:drawing")):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        nxt = body[i + 1]
        if nxt.text and len(nxt.text) < 160 and not nxt.style.name.startswith("Heading"):
            for r in nxt.runs: r.font.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = MUTED

def add_field(par, instr):
    r = par.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
    r2 = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r3._r.append(f2)
    for x in (r, r2, r3): x.font.size = Pt(8); x.font.color.rgb = MUTED; x.font.name = "Arial"

for s in d.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2.2)
    hp = s.header.paragraphs[0] if s.header.paragraphs else s.header.add_paragraph()
    hp.text = ""; rr = hp.add_run("Rikolto RDC · Rapport à la direction · Café & Cacao 2019-2025")
    rr.font.size = Pt(8); rr.font.color.rgb = MUTED; rr.font.name = "Arial"; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
    fp.text = ""; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a = fp.add_run("Page "); a.font.size = Pt(8); a.font.color.rgb = MUTED; a.font.name = "Arial"
    add_field(fp, "PAGE")
    b = fp.add_run(" sur "); b.font.size = Pt(8); b.font.color.rgb = MUTED; b.font.name = "Arial"
    add_field(fp, "NUMPAGES")

p1 = d.paragraphs[1]
for r in p1.runs:
    r.text = r.text.replace("Sep 30, 2026", "30 septembre 2026").replace("@Lucien", "Lucien")
cp = d.core_properties
cp.author = "Lucien Buzera"; cp.title = "Rapport à la direction – Café & Cacao 2019-2025"; cp.language = "fr-FR"; cp.subject = "Rikolto RDC"
d.save(OUT)
print(OUT)
